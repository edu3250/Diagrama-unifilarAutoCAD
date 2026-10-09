"""Stage 3.6.1: calculation report (xlsx + pdf), project sheet and bill of materials."""

from __future__ import annotations

from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from pvsld.cli import main
from pvsld.core.validation import ValidationReport, validate_pv_design
from pvsld.report import (
    bill_of_materials,
    bom_text,
    build_memoria,
    write_memoria_pdf,
    write_memoria_xlsx,
    write_project_sheet,
)
from sizing_helpers import ET_550, FIXTURE_RECORDS, registry, size


@pytest.fixture(scope="module")
def design() -> tuple[ValidationReport, str]:
    result = size(ET_550, "auto", module_count_min=5, module_count_max=5)
    assert result.selected is not None
    assert result.spec is not None
    report = validate_pv_design(result.spec)
    assert report.ok
    assert report.spec is not None
    assert report.derived is not None
    return report, result.selected.explanation_es


def test_the_memoria_matches_the_derived_values_and_every_check_passes(
    design: tuple[ValidationReport, str],
) -> None:
    report, why = design
    assert report.spec is not None
    assert report.derived is not None
    mem = build_memoria(report.spec, report.derived, report.findings, explanation_es=why)
    assert mem.ok
    values = mem.values
    circuits = {c.circuit_id: c for c in report.derived.circuits}
    dc = next(c for c in report.spec.circuits if c.kind == "pv_source")
    ac = next(c for c in report.spec.circuits if c.kind == "inverter_output")
    assert values["vddc"] == pytest.approx(circuits[dc.id].vd_pct, rel=1e-6)
    assert values["vdac"] == pytest.approx(circuits[ac.id].vd_pct, rel=1e-6)
    assert values["vocstr"] == pytest.approx(report.derived.strings[0].voc_max_string_v, rel=1e-6)
    sections = list(dict.fromkeys(c.section for c in mem.calcs))
    assert sections[0] == "2. Arreglo"
    assert "8. Canalización" in sections  # EMT raceways from the sizing engine
    assert {"protchk", "swchk", "kaicchk", "rule120", "fillDC", "fillAC"} <= set(values)


def test_the_workbook_has_live_formulas_over_the_inputs(
    design: tuple[ValidationReport, str], tmp_path: Path
) -> None:
    report, why = design
    assert report.spec is not None
    assert report.derived is not None
    mem = build_memoria(report.spec, report.derived, report.findings, explanation_es=why)
    path = write_memoria_xlsx(mem, tmp_path / "m.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Datos", "Memoria", "Materiales", "Validación"]
    memoria = wb["Memoria"]
    formulas = [
        c.value
        for row in memoria.iter_rows(min_col=3, max_col=3)
        for c in row
        if isinstance(c.value, str) and c.value.startswith("=")
    ]
    assert len(formulas) == len(mem.calcs)
    assert all("Datos!$C$" in f or "Memoria!$C$" in f for f in formulas)
    checks = [
        c.value
        for row in memoria.iter_rows(min_col=6, max_col=6)
        for c in row
        if isinstance(c.value, str) and c.value.startswith("=IF(")
    ]
    assert len(checks) == sum(1 for c in mem.calcs if c.check is not None)


def test_the_pdf_is_written(design: tuple[ValidationReport, str], tmp_path: Path) -> None:
    report, why = design
    assert report.spec is not None
    assert report.derived is not None
    mem = build_memoria(report.spec, report.derived, report.findings, explanation_es=why)
    data = write_memoria_pdf(mem, tmp_path / "m.pdf").read_bytes()
    assert data.startswith(b"%PDF")
    assert data.count(b"/Type /Page") >= 3  # pages and the page tree


def test_the_bill_of_materials_names_the_catalogue_devices(
    design: tuple[ValidationReport, str],
) -> None:
    report, _ = design
    assert report.spec is not None
    assert report.derived is not None
    lines = bill_of_materials(report.spec, report.derived)
    text = bom_text(report.spec, report.derived, lines)
    assert lines[0].item == "Módulo fotovoltaico"
    assert lines[0].quantity == 5
    assert "Littelfuse SPF0" in text  # the gPV fuse of the box
    assert "Suntree SISO-40-32 4 polos" in text
    assert "Square D QO2100" in text
    assert text.splitlines()[0].startswith("RESUMEN DE LA INSTALACIÓN")


def test_the_project_sheet_is_prefilled_with_catalogue_dropdowns(
    design: tuple[ValidationReport, str], tmp_path: Path
) -> None:
    report, why = design
    path = write_project_sheet(
        tmp_path / "h.xlsx", registry(), report.spec, report.derived, explanation_es=why
    )
    wb = load_workbook(path)
    assert wb.sheetnames == [
        "Instrucciones",
        "Proyecto",
        "Equipos",
        "Protecciones",
        "Conductores",
        "Tierra",
        "Catálogo",
    ]
    assert wb["Catálogo"].sheet_state == "hidden"
    equipos = wb["Equipos"]
    assert equipos["B6"].value.startswith("ETSOLAR-")  # the module's catalogue id
    assert equipos["B12"].value == "AUTO"  # inverter
    assert any(dv.formula1 == "=L_Inversores" for dv in equipos.data_validations.dataValidation)
    assert "L_Fusibles" in wb.defined_names


def test_the_report_command_writes_the_four_files(
    design: tuple[ValidationReport, str], tmp_path: Path
) -> None:
    report, _ = design
    assert report.spec is not None
    spec_file = tmp_path / "spec.yaml"
    spec_file.write_text(
        yaml.safe_dump(report.spec.model_dump(mode="json", by_alias=True), allow_unicode=True),
        encoding="utf-8",
    )
    out = tmp_path / "out"
    assert (
        main(["report", str(spec_file), "-o", str(out), "--catalogue", str(FIXTURE_RECORDS)]) == 0
    )
    assert {p.name for p in out.iterdir()} == {
        "memoria_calculo.xlsx",
        "memoria_calculo.pdf",
        "hoja_de_proyecto.xlsx",
        "resumen_materiales.txt",
    }


def test_the_memoria_shows_the_personal_data_given_and_blanks_the_rest(
    design: tuple[ValidationReport, str],
) -> None:
    report, _ = design
    assert report.spec is not None
    assert report.derived is not None
    general = dict(build_memoria(report.spec, report.derived).general)
    assert general["Propietario"] == "Juan Pérez, RPU 000000000000"
    assert general["Ubicación"].startswith("Av. Ejemplo 123, Centro, Zapopan, Jalisco, C.P. 45000")
    assert "Ing. Nombre Apellido" in general["Responsable"]
    project = report.spec.project.model_copy(
        update={"client": report.spec.project.client.model_construct()}
    )
    blank_utility = report.spec.utility.model_copy(update={"rpu": None})
    anonymous = report.spec.model_copy(update={"project": project, "utility": blank_utility})
    general = dict(build_memoria(anonymous, report.derived).general)
    assert general["Propietario"] == ""
