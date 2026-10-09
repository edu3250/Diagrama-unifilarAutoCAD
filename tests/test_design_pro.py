"""Stage 3.6.2b: the professional mode, driven by the project sheet the user fills in.

A quick design on the fixture catalogue writes the prefilled sheet (8 x ETSOLAR 550 W); each test
edits a copy of it as a user would in Excel and designs from it, without AutoCAD.
"""

from __future__ import annotations

import shutil
from datetime import date
from pathlib import Path
from typing import Any

import pytest
import yaml
from mcp.server.mcpserver.exceptions import ToolError
from openpyxl import Workbook, load_workbook

from pvsld.cli import main
from pvsld.design import design_and_draw
from pvsld.design.pipeline import PROJECT_SHEET
from pvsld.design.pro import (
    ProjectSheetError,
    design_from_sheet,
    read_project_sheet,
    sheet_request,
    write_reviewed_sheet,
)
from pvsld.finishers.core_console import AutocadInfo
from pvsld.mcp.design_tool import REVIEW_SHEET, run_new_sheet, run_pro_design
from pvsld.mcp.sandbox import OutputSandbox
from sizing_helpers import ET_550, FIXTURE_RECORDS, registry

TODAY = date(2026, 10, 9)
NO_AUTOCAD = AutocadInfo(None)
QUICK = {"module": ET_550, "module_count_min": 8, "module_count_max": 8}


@pytest.fixture(scope="module")
def prefilled(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """The project sheet a quick design leaves in its folder."""
    folder = tmp_path_factory.mktemp("quick")
    result = design_and_draw(
        QUICK, registry(), folder, use_autocad="never", autocad=NO_AUTOCAD, today=TODAY
    )
    assert result.ok
    return result.files[PROJECT_SHEET]


@pytest.fixture
def sheet(prefilled: Path, tmp_path: Path) -> Path:
    copy = tmp_path / PROJECT_SHEET
    shutil.copy(prefilled, copy)
    return copy


def _row(ws: Any, label: str) -> int:
    """The parameter row of ``label`` (a section title of the same name has no options)."""
    return next(
        r for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value == label and ws.cell(r, 4).value
    )


def edit(path: Path, changes: dict[tuple[str, str], Any]) -> None:
    """Type values into column B as a user would."""
    workbook = load_workbook(path)
    for (name, label), value in changes.items():
        ws = workbook[name]
        ws.cell(_row(ws, label), 2).value = value
    workbook.save(path)


def design(sheet: Path, folder: Path) -> Any:
    return design_from_sheet(
        sheet, registry(), folder, use_autocad="never", autocad=NO_AUTOCAD, today=TODAY
    )


# --- reading -------------------------------------------------------------------------------------


def test_the_sheet_is_read_by_label(sheet: Path) -> None:
    cells = read_project_sheet(sheet)
    assert cells[("Equipos", "Módulo fotovoltaico")].value == ET_550
    assert cells[("Equipos", "Cantidad total de módulos")].value == 8
    assert cells[("Equipos", "Inversor")].value == "AUTO"
    assert cells[("Proyecto", "RPU")].row > 5


def test_moved_rows_are_still_found(sheet: Path) -> None:
    workbook = load_workbook(sheet)
    workbook["Equipos"].insert_rows(6, amount=3)
    workbook.save(sheet)
    request, issues = sheet_request(read_project_sheet(sheet), registry(), today=TODAY)
    assert issues == []
    assert request["module"] == ET_550


def test_the_prefilled_sheet_asks_for_the_quick_design(sheet: Path) -> None:
    request, issues = sheet_request(read_project_sheet(sheet), registry(), today=TODAY)
    assert issues == []
    assert (request["module"], request["module_count_min"]) == (ET_550, 8)
    assert request["inverters"] == "auto"
    assert "n_strings" not in request
    assert "dc_fuses" not in request
    assert request["template"]["project"]["client"] == {}
    assert request["template"]["utility"]["system"] == "2F-3H"


@pytest.mark.parametrize("content", [b"not an excel file", None])
def test_a_file_that_is_not_a_project_sheet_is_refused(tmp_path: Path, content: Any) -> None:
    path = tmp_path / "otra.xlsx"
    if content is None:
        Workbook().save(path)
    else:
        path.write_bytes(content)
    with pytest.raises(ProjectSheetError):
        read_project_sheet(path)


# --- designing -----------------------------------------------------------------------------------


def test_the_untouched_sheet_designs_and_keeps_the_users_workbook(
    sheet: Path, tmp_path: Path
) -> None:
    result = design(sheet, tmp_path / "out")
    assert result.ok
    assert result.issues == ()
    returned = load_workbook(result.design.files[PROJECT_SHEET])["Equipos"]
    row = _row(returned, "Inversor")
    assert returned.cell(row, 2).value == "AUTO"  # the user's choice stays
    assert returned.cell(row, 5).value  # the last calculation is refreshed
    assert not list((tmp_path / "out").glob(".*.xlsx"))  # the fresh copy is removed


def test_fixed_values_and_personal_data_reach_the_design(sheet: Path, tmp_path: Path) -> None:
    edit(
        sheet,
        {
            ("Equipos", "Número de cadenas"): 1,
            ("Equipos", "Módulos por cadena"): 8,
            ("Conductores", "Calibre de CD"): "6 AWG",
            ("Proyecto", "Titular / propietario"): "Ana López",
            ("Proyecto", "RPU"): "1234 5678 9012",
            ("Proyecto", "Código postal"): "45000",
            ("Proyecto", "Empresa instaladora"): "Solar Bajío",
            ("Proyecto", "Fecha"): "01/02/2026",
        },
    )
    result = design(sheet, tmp_path / "out")
    assert result.ok, result.issues
    spec = yaml.safe_load(result.design.files["especificacion.yaml"].read_text(encoding="utf-8"))
    assert [s["n_series"] for s in spec["strings"]] == [8]
    assert spec["circuits"][0]["conductors"]["size"] == "6 AWG"
    assert spec["project"]["client"] == {"name": "Ana López"}
    assert spec["utility"]["rpu"] == "123456789012"
    assert spec["project"]["site"]["address"] == {"cp": "45000"}
    assert spec["title_block"]["responsible"] == {"company": "Solar Bajío"}
    assert spec["title_block"]["date"] == date(2026, 2, 1)


def test_unusable_cells_are_reported_on_their_cells(sheet: Path, tmp_path: Path) -> None:
    edit(
        sheet,
        {
            ("Proyecto", "Código postal"): "4500",
            ("Proyecto", "RPU"): "123",
            ("Proyecto", "Cédula profesional"): "12",
            ("Proyecto", "Tensión nominal"): 230,
            ("Proyecto", "Sistema de suministro"): "3F-4H",
            ("Equipos", "Número de cadenas"): 3,
            ("Equipos", "Módulos por cadena"): 3,
            ("Equipos", "Inversor"): "FOO-INVERSOR-1",
            ("Conductores", "Longitud inversor → ITM-1"): "treinta",
        },
    )
    result = design(sheet, tmp_path / "out")
    assert not result.ok
    assert result.stage == "sheet"
    assert result.design is None
    labels = {issue.label for issue in result.issues}
    assert labels == {
        "Código postal",
        "RPU",
        "Cédula profesional",
        "Tensión nominal",
        "Sistema de suministro",
        "Número de cadenas",
        "Módulos por cadena",
        "Inversor",
        "Longitud inversor → ITM-1",
    }
    assert result.missing_components == ("FOO-INVERSOR-1",)
    assert all(issue.row for issue in result.issues)
    assert not (tmp_path / "out").exists()


def test_a_fixed_value_the_calculation_rejects_is_put_back_on_its_cell(
    sheet: Path, tmp_path: Path
) -> None:
    edit(sheet, {("Protecciones", "Fusible gPV"): "LITTELFUSE-SPF015"})
    result = design(sheet, tmp_path / "out")
    assert result.stage == "sizing"
    (issue,) = result.issues
    assert (issue.sheet, issue.label) == ("Protecciones", "Fusible gPV")
    assert issue.message_es.startswith("OCP-002: Ningún fusible gPV del catálogo cumple")
    assert "SPF015: 15 A" in issue.message_es


def test_the_reviewed_sheet_marks_the_cells_in_red(sheet: Path, tmp_path: Path) -> None:
    edit(sheet, {("Proyecto", "Código postal"): "4500"})
    result = design(sheet, tmp_path / "out")
    reviewed = write_reviewed_sheet(sheet, tmp_path / REVIEW_SHEET, result.issues)
    workbook = load_workbook(reviewed)
    (issue,) = result.issues
    cell = workbook["Proyecto"].cell(issue.row, 2)
    assert cell.value == "4500"  # the user's value is kept
    assert cell.fill.fgColor.rgb.endswith("FFC7CE")
    assert cell.comment is not None
    assert "código postal" in cell.comment.text
    guide = workbook["Instrucciones"]
    assert any(
        str(guide.cell(r, 2).value).startswith("REVISIÓN: 1 dato")
        for r in range(1, guide.max_row + 1)
    )


# --- MCP tools -----------------------------------------------------------------------------------


def _pro(box: OutputSandbox, name: str) -> Any:
    return run_pro_design(
        box,
        registry(),
        name=name,
        use_autocad="auto",
        max_preview_bytes=62_000,
        autocad=NO_AUTOCAD,
        today=TODAY,
    )


def test_new_project_sheet_writes_a_blank_sheet_once(tmp_path: Path) -> None:
    box = OutputSandbox.at(tmp_path)
    output = run_new_sheet(box, registry(), name="casa", overwrite=False)
    assert output.ok
    cells = read_project_sheet(tmp_path / "casa" / PROJECT_SHEET)
    assert cells[("Equipos", "Inversor")].value == "AUTO"
    with pytest.raises(ToolError, match="overwrite=true"):
        run_new_sheet(box, registry(), name="casa", overwrite=False)


def test_the_pro_tool_reports_cells_then_designs_once_they_are_fixed(
    prefilled: Path, tmp_path: Path
) -> None:
    box = OutputSandbox.at(tmp_path)
    folder = tmp_path / "casa"
    folder.mkdir()
    shutil.copy(prefilled, folder / PROJECT_SHEET)
    edit(folder / PROJECT_SHEET, {("Equipos", "Inversor"): "FOO-INVERSOR-1"})

    output, image = _pro(box, "casa")
    assert not output.ok
    assert image is None
    assert output.missing_components == ["FOO-INVERSOR-1"]
    assert output.issues[0]["label"] == "Inversor"
    assert (folder / REVIEW_SHEET).is_file()
    assert not (folder / "unifilar.dxf").exists()

    edit(folder / PROJECT_SHEET, {("Equipos", "Inversor"): "AUTO"})
    output, image = _pro(box, "casa")
    assert output.ok
    assert image is not None
    assert (folder / "unifilar.dxf").is_file()
    assert not (folder / REVIEW_SHEET).exists()


def test_the_pro_tool_needs_a_sheet_in_the_folder(tmp_path: Path) -> None:
    with pytest.raises(ToolError, match=r"no hoja_de_proyecto\.xlsx"):
        _pro(OutputSandbox.at(tmp_path), "vacio")


# --- CLI ------------------------------------------------------------------------------------------


def test_pvsld_pro_new_and_design(
    sheet: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    catalogue = ["--catalogue", str(FIXTURE_RECORDS)]
    assert main(["pro", "new", str(tmp_path / "nuevo"), *catalogue]) == 0
    assert (tmp_path / "nuevo" / PROJECT_SHEET).is_file()

    edit(sheet, {("Proyecto", "Código postal"): "4500"})
    argv = ["pro", "design", str(sheet), "-o", str(tmp_path / "p"), "--no-autocad", *catalogue]
    assert main(argv) == 1
    assert "Proyecto > Código postal" in capsys.readouterr().err
    assert sheet.with_name("hoja_de_proyecto_revisar.xlsx").is_file()

    edit(sheet, {("Proyecto", "Código postal"): "45000"})
    assert main(argv) == 0
    assert capsys.readouterr().out.startswith("why: Se eligió")


def test_sheet_and_calculation_problems_come_in_one_round(sheet: Path, tmp_path: Path) -> None:
    edit(
        sheet,
        {
            ("Proyecto", "Código postal"): "4500",
            ("Protecciones", "Fusible gPV"): "LITTELFUSE-SPF015",
        },
    )
    result = design(sheet, tmp_path / "out")
    assert result.stage == "sheet"
    assert [issue.label for issue in result.issues] == ["Código postal", "Fusible gPV"]
