"""Professional mode (Stage 3.6.2b): design an installation from the project sheet the user filled.

The project sheet (:func:`pvsld.report.write_project_sheet`) has one row per parameter: the label
in column A, the user's value in column B (AUTO lets the calculation decide) and the last
calculation in column E. :func:`read_project_sheet` reads the rows by label, so a user who inserts
or moves rows does not break it; :func:`sheet_request` checks every value and turns the sheet into a
sizing request whose fixed values the engine keeps (Stage 3.6.2a); :func:`design_from_sheet` runs
the one-step design with it.

A value that cannot be used is a :class:`SheetIssue` on its cell, and so is a fixed value the
calculation rejects (a fuse below 1.56 × Isc, a conductor too small): :func:`write_reviewed_sheet`
returns the user's own workbook with those cells in red and the reason in a comment. A model that
is not in the catalogue is reported as a missing component, so the plugin can offer to add it from
its datasheet (Stage 3.6.3).

Personal data goes to the drawing and the report only when the user fills it in.
"""

from __future__ import annotations

import re
from collections.abc import Callable, Iterable, Mapping
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Literal

from pvsld.catalogue.registry import ComponentRegistry
from pvsld.core.model import AWG_PATTERN, NOMINAL_VOLTAGES_V
from pvsld.core.tables import get_tables
from pvsld.design.pipeline import DesignResult, design_and_draw, load_quick_template
from pvsld.finishers.core_console import AutocadInfo, Runner, run_process
from pvsld.sizing import SizingInputError, request_from_mapping, size_pv_system
from pvsld.sizing.models import SUPPORTED_SYSTEMS, SizingResult

__all__ = [
    "INPUT_SHEETS",
    "ProResult",
    "ProjectSheetError",
    "SheetCell",
    "SheetIssue",
    "design_from_sheet",
    "read_project_sheet",
    "refresh_last_calculation",
    "rejection_issues",
    "sheet_request",
    "write_reviewed_sheet",
]

INPUT_SHEETS = ("Proyecto", "Equipos", "Protecciones", "Conductores", "Tierra")
FIRST_ROW = 5
AUTO = "AUTO"
REVIEW_FILL = "FFC7CE"
REGIMES = {
    "medición neta": "medicion_neta",
    "facturación neta": "facturacion_neta",
    "venta total": "venta_total",
}
DC_SYSTEMS = {"no aterrizado (690-35)": "ungrounded_690_35", "aterrizado": "grounded"}
DEFAULT_ELECTRODE_LENGTH_M = 3.05

Key = tuple[str, str]
"""(sheet, label) of a row."""


class ProjectSheetError(ValueError):
    """The file cannot be read as a project sheet; the message is safe to show the user."""


@dataclass(frozen=True)
class SheetCell:
    """One parameter row: where it is and what the user wrote in column B."""

    sheet: str
    row: int
    label: str
    value: Any


@dataclass(frozen=True)
class SheetIssue:
    """A value that cannot be used, on its cell; ``missing_component`` names an unknown model."""

    sheet: str
    label: str
    message_es: str
    row: int | None = None
    missing_component: str | None = None

    def to_dict(self) -> dict[str, Any]:
        data: dict[str, Any] = {
            "sheet": self.sheet,
            "label": self.label,
            "row": self.row,
            "message_es": self.message_es,
        }
        if self.missing_component:
            data["missing_component"] = self.missing_component
        return data


def read_project_sheet(path: Path) -> dict[Key, SheetCell]:
    """Every parameter row of the input sheets, keyed by (sheet, label).

    Raises:
        ProjectSheetError: the file is missing, is not an Excel workbook or lacks an input sheet.
    """
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except FileNotFoundError:
        raise ProjectSheetError(f"no existe el archivo {path.name}") from None
    except Exception as error:  # openpyxl raises several unrelated types for a bad file
        raise ProjectSheetError(f"{path.name} no es un libro de Excel válido: {error}") from error
    try:
        missing = [name for name in INPUT_SHEETS if name not in workbook.sheetnames]
        if missing:
            raise ProjectSheetError(
                f"{path.name} no es una hoja de proyecto: le faltan las hojas {', '.join(missing)}"
            )
        cells: dict[Key, SheetCell] = {}
        for name in INPUT_SHEETS:
            rows = workbook[name].iter_rows(min_row=FIRST_ROW, max_col=4, values_only=True)
            for row, (label, value, _unit, options) in enumerate(rows, start=FIRST_ROW):
                # A section title ("Inversor") has a label but no options column.
                if isinstance(label, str) and label.strip() and options:
                    text = value.strip() if isinstance(value, str) else value
                    cells[(name, label.strip())] = SheetCell(name, row, label.strip(), text)
        return cells
    finally:
        workbook.close()


# --- reading values ------------------------------------------------------------------------------


class _Values:
    """Typed access to the cells; every problem becomes a :class:`SheetIssue`."""

    def __init__(self, cells: Mapping[Key, SheetCell], registry: ComponentRegistry) -> None:
        self.cells = cells
        self.registry = registry
        self.issues: list[SheetIssue] = []

    def issue(self, key: Key, message: str, missing: str | None = None) -> None:
        cell = self.cells.get(key)
        self.issues.append(SheetIssue(key[0], key[1], message, cell.row if cell else None, missing))

    def raw(self, key: Key) -> Any:
        cell = self.cells.get(key)
        if cell is None:
            self.issues.append(
                SheetIssue(key[0], key[1], f"Falta el renglón «{key[1]}» en la hoja {key[0]}.")
            )
            return None
        value = cell.value
        return None if value == "" else value

    def text(self, key: Key) -> str | None:
        value = self.raw(key)
        if value is None:
            return None
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        return str(value).strip() or None

    def auto(self, key: Key) -> bool:
        value = self.text(key)
        return value is None or value.upper() == AUTO

    def number(
        self,
        key: Key,
        *,
        low: float | None = None,
        high: float | None = None,
        integer: bool = False,
        required: bool = True,
    ) -> float | None:
        value = self.raw(key)
        if value is None:
            if required:
                self.issue(key, "Falta el valor.")
            return None
        try:
            number = float(str(value).replace(",", ".")) if isinstance(value, str) else float(value)
        except (TypeError, ValueError):
            self.issue(key, f"«{value}» no es un número.")
            return None
        if integer and not number.is_integer():
            self.issue(key, f"{number:g} debe ser un número entero.")
            return None
        if (low is not None and number < low) or (high is not None and number > high):
            bounds = (
                f"entre {low:g} y {high:g}"
                if low is not None and high is not None
                else (f"de al menos {low:g}" if low is not None else f"de máximo {high:g}")
            )
            self.issue(key, f"{number:g} está fuera de rango: debe ser {bounds}.")
            return None
        return number

    def pattern(self, key: Key, regex: str, what: str) -> str | None:
        text = self.text(key)
        if text is None:
            return None
        compact = re.sub(r"[\s-]", "", text) if what != "correo" else text
        if not re.fullmatch(regex, compact):
            self.issue(key, f"«{text}» no es {what} válido.")
            return None
        return compact

    def choice(self, key: Key, options: Iterable[str], *, default: str) -> str:
        text = self.text(key)
        if text is None:
            return default
        for option in options:
            if option.casefold() == text.casefold():
                return option
        self.issue(key, f"«{text}» no es una opción de la lista: {', '.join(options)}.")
        return default

    def component(
        self, key: Key, ids: Iterable[str], what: str, *, auto: bool = True
    ) -> str | None:
        """A catalogue id, ``None`` for AUTO; an unknown model is a missing component."""
        text = self.text(key)
        if text is None or (auto and text.upper() == AUTO):
            return None
        known = {i.upper(): i for i in ids}
        if text.upper() in known:
            return known[text.upper()]
        self.issue(
            key,
            f"El {what} «{text}» no está en el catálogo. Agréguelo desde su hoja técnica "
            "(/unifilar-catalogo) o elija uno de la lista.",
            missing=text,
        )
        return None

    def size(self, key: Key) -> str | None:
        text = self.text(key)
        if text is None or text.upper() == AUTO:
            return None
        normal = re.sub(r"\s*awg$", " AWG", text, flags=re.IGNORECASE)
        if not re.fullmatch(AWG_PATTERN, normal):
            self.issue(key, f"«{text}» no es un calibre AWG de la lista (por ejemplo 10 AWG).")
            return None
        return normal

    def day(self, key: Key) -> date | None:
        value = self.raw(key)
        if value is None:
            return None
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        try:
            return datetime.strptime(str(value).strip(), "%d/%m/%Y").date()
        except ValueError:
            self.issue(key, f"«{value}» no es una fecha dd/mm/aaaa.")
            return None


def _emt_trade_size(values: _Values, key: Key, designations: Iterable[float]) -> float | None:
    text = values.text(key)
    if text is None or text.upper() == AUTO:
        return None
    match = re.match(r"\s*(\d+(?:[.,]\d+)?)", text)
    size = float(match.group(1).replace(",", ".")) if match else None
    if size not in set(designations):
        values.issue(key, f'«{text}» no es un tubo EMT de la lista (por ejemplo 21 mm (3/4")).')
        return None
    return size


def _electrode(text: str) -> dict[str, Any]:
    """ "Varilla copperweld 3.05 m" -> type and length; other electrodes keep the default length."""
    match = re.fullmatch(r"(.*?)\s+(\d+(?:\.\d+)?)\s*m", text)
    if match:
        return {"type": match.group(1).lower(), "length_m": float(match.group(2))}
    return {"type": text.lower(), "length_m": DEFAULT_ELECTRODE_LENGTH_M}


def sheet_request(
    cells: Mapping[Key, SheetCell], registry: ComponentRegistry, *, today: date | None = None
) -> tuple[dict[str, Any], list[SheetIssue]]:
    """The sizing request (with its template) the sheet describes, and every unusable value."""
    v = _Values(cells, registry)
    template = load_quick_template()
    project, site = template["project"], template["project"]["site"]
    utility, ac_bos, dc_bos = template["utility"], template["ac_bos"], template["dc_bos"]
    block, grounding = template["title_block"], template["grounding"]
    request: dict[str, Any] = {"template": template, "dc_ocpd": "always"}

    def put(target: dict[str, Any], name: str, value: Any) -> None:
        if value is not None:
            target[name] = value

    # Proyecto: personal data only when given; service; responsible.
    p = "Proyecto"
    put(project, "name", v.text((p, "Nombre del proyecto")))
    client = {"name": v.text((p, "Titular / propietario"))}
    project["client"] = {k: x for k, x in client.items() if x is not None}
    address = {
        "street": v.text((p, "Calle y número")),
        "colonia": v.text((p, "Colonia")),
        "municipio": v.text((p, "Municipio")),
        "estado": v.text((p, "Estado")),
        "cp": v.pattern((p, "Código postal"), r"\d{5}", "un código postal de 5 dígitos"),
    }
    site["address"] = {k: x for k, x in address.items() if x is not None}
    put(site, "lat", v.number((p, "Latitud"), low=14, high=33, required=False))
    put(site, "lon", v.number((p, "Longitud"), low=-118, high=-86, required=False))
    put(utility, "rpu", v.pattern((p, "RPU"), r"\d{12}", "un RPU de 12 dígitos"))
    put(utility, "service_number", v.text((p, "Número de servicio")))
    meter = v.text((p, "Número de medidor"))
    put(utility, "meter_number", meter)
    if meter:
        ac_bos["meters"][0]["meter_no"] = meter
    put(utility, "tariff", v.text((p, "Tarifa")))
    system = v.choice((p, "Sistema de suministro"), ["1F-2H", "2F-3H", "3F-4H"], default="2F-3H")
    if system not in SUPPORTED_SYSTEMS:
        v.issue(
            (p, "Sistema de suministro"),
            f"Esta versión diseña servicios {' y '.join(SUPPORTED_SYSTEMS)}; {system} todavía no.",
        )
    else:
        utility["system"] = ac_bos["panels"][0]["system"] = system
    voltage = v.number((p, "Tensión nominal"))
    if voltage is not None:
        if int(voltage) not in NOMINAL_VOLTAGES_V:
            v.issue((p, "Tensión nominal"), f"{voltage:g} V no es una tensión normalizada de CFE.")
        else:
            utility["nominal_voltage_v"] = int(voltage)
    put(
        utility, "available_fault_current_ka", v.number((p, "Corriente de falla disponible"), low=1)
    )
    regime = v.choice(
        (p, "Régimen"),
        ["Medición neta", "Facturación neta", "Venta total"],
        default="Medición neta",
    )
    utility["regime"] = REGIMES[regime.casefold()]
    responsible = {
        "name": v.text((p, "Ingeniero responsable")),
        "cedula_profesional": v.pattern(
            (p, "Cédula profesional"), r"\d{7,8}", "una cédula profesional de 7 u 8 dígitos"
        ),
        "company": v.text((p, "Empresa instaladora")),
    }
    block["responsible"] = {k: x for k, x in responsible.items() if x is not None}
    put(block, "uvie", v.text((p, "UVIE")))
    day = v.day((p, "Fecha")) or today or date.today()
    block["date"] = day
    for revision in block["revisions"]:
        revision["date"] = day

    # Equipos: module, counts, inverter.
    e = "Equipos"
    module = v.component(
        (e, "Módulo fotovoltaico"),
        (c.component_id for c in registry.modules()),
        "módulo",
        auto=False,
    )
    if module is None and v.text((e, "Módulo fotovoltaico")) is None:
        v.issue((e, "Módulo fotovoltaico"), "Elija el módulo fotovoltaico.")
    request["module"] = module
    total = v.number((e, "Cantidad total de módulos"), low=1, high=200, integer=True)
    if total is not None:
        request["module_count_min"] = request["module_count_max"] = int(total)
    strings = (
        None
        if v.auto((e, "Número de cadenas"))
        else v.number((e, "Número de cadenas"), low=1, high=12, integer=True)
    )
    series = (
        None
        if v.auto((e, "Módulos por cadena"))
        else v.number((e, "Módulos por cadena"), low=1, high=40, integer=True)
    )
    if strings is not None:
        request["n_strings"] = int(strings)
    if series is not None:
        request["n_series"] = int(series)
    if total is not None:
        if strings is not None and series is not None and strings * series != total:
            v.issue(
                (e, "Módulos por cadena"),
                f"{strings:g} cadenas × {series:g} módulos = {strings * series:g}, no los "
                f"{total:g} módulos de la cantidad total.",
            )
        for fixed, label in ((strings, "Número de cadenas"), (series, "Módulos por cadena")):
            if fixed is not None and total % fixed:
                v.issue(
                    (e, label), f"{total:g} módulos no se reparten en partes iguales con {fixed:g}."
                )
    request["optimizers_on_all_modules"] = (
        v.choice((e, "Optimizadores en todos los módulos"), ["No", "Sí"], default="No") == "Sí"
    )
    inverter = v.component(
        (e, "Inversor"), (c.component_id for c in registry.inverters()), "inversor"
    )
    request["inverters"] = [inverter] if inverter else "auto"

    # Protecciones.
    r = "Protecciones"
    device = v.choice(
        (r, "Protección de cadena"), ["Fusible gPV", "ITM de CD"], default="Fusible gPV"
    )
    request["dc_ocpd_device"] = "fuse" if device == "Fusible gPV" else "breaker"
    fuse = v.component((r, "Fusible gPV"), (c.component_id for c in registry.dc_fuses()), "fusible")
    if fuse and device == "Fusible gPV":
        request["dc_fuses"], request["dc_fuse_fallback"] = [fuse], False
    breaker = v.component(
        (r, "ITM de CD (si elige ITM)"),
        (c.component_id for c in registry.dc_breakers()),
        "ITM de CD",
    )
    if breaker:
        request["dc_breakers"] = [breaker]
    switch = v.component(
        (r, "Seccionador de la caja"),
        (c.component_id for c in registry.dc_switches()),
        "seccionador",
    )
    if switch:
        request["dc_switches"] = [switch]
    dc_spd = dc_bos["spds"][0]
    dc_spd["spd_type"] = v.choice((r, "DPS de CD: tipo"), ["T2", "T1+T2"], default="T2")
    put(dc_spd, "in_ka", v.number((r, "DPS de CD: corriente nominal In"), low=1))
    ac_breakers = list(registry.ac_breakers())
    itm1 = v.component(
        (r, "ITM-1 (salida del inversor)"), (c.component_id for c in ac_breakers), "ITM"
    )
    if itm1:
        request["ac_breakers"] = [itm1]
    main = v.component(
        (r, "ITM-P (principal del servicio)"),
        (c.component_id for c in ac_breakers),
        "ITM",
        auto=False,
    )
    if main:
        request["main_breaker"] = main
        rating = next(c.rated_current_a for c in ac_breakers if c.component_id == main)
        ac_bos["main_breakers"][0]["rating_a"] = rating
        utility["service_main"]["rating_a"] = rating
    put(ac_bos["panels"][0], "bus_a", v.number((r, "Barra del centro de carga"), low=1))
    ac_spd = ac_bos["spds"][0]
    ac_spd["spd_type"] = v.choice((r, "DPS de CA: tipo"), ["T2", "T1+T2"], default="T2")
    put(ac_spd, "uc_v", v.number((r, "DPS de CA: Uc"), low=1))

    # Conductores.
    c = "Conductores"
    routing: dict[str, Any] = {}
    put(
        routing,
        "dc_string_length_m",
        v.number((c, "Longitud cadenas → caja → inversor"), low=0.1, high=999),
    )
    put(
        routing, "ac_output_length_m", v.number((c, "Longitud inversor → ITM-1"), low=0.1, high=999)
    )
    put(
        routing,
        "rooftop_clearance_mm",
        v.number((c, "Separación del tubo sobre el techo"), low=0, high=1000),
    )
    emt = [d for d, _t, _diameter, _area in get_tables(template["standards"]["nom_edition"]).emt]
    put(routing, "dc_raceway_trade_size_mm", _emt_trade_size(v, (c, "Tubo de CD (EMT)"), emt))
    put(routing, "ac_raceway_trade_size_mm", _emt_trade_size(v, (c, "Tubo de CA (EMT)"), emt))
    request["routing"] = routing
    cables = {cable.component_id: cable for cable in registry.cables()}
    cable = v.component((c, "Cable de CD"), cables, "cable")
    dc_size = v.size((c, "Calibre de CD"))
    if cable:
        request["dc_cable"] = cables[cable].family_id
        dc_size = dc_size or cables[cable].size  # a cable reference names its size
    put(request, "dc_conductor_size", dc_size)
    put(request, "ac_conductor_size", v.size((c, "Calibre de CA")))
    limits = template["standards"]["vd_limits_pct"]
    put(limits, "dc", v.number((c, "Caída de tensión máxima CD"), low=0.1, high=5))
    put(limits, "ac", v.number((c, "Caída de tensión máxima CA"), low=0.1, high=5))

    # Tierra.
    t = "Tierra"
    kind = v.text((t, "Tipo de electrodo"))
    electrode = grounding["electrodes"][0]
    if kind:
        electrode.update(_electrode(kind))
    qty = v.number((t, "Cantidad de electrodos"), low=1, high=20, integer=True)
    if qty is not None:
        electrode["qty"] = int(qty)
    put(grounding, "design_resistance_ohm", v.number((t, "Resistencia de diseño"), low=0.1))
    put(grounding, "gec_dc", v.size((t, "Conductor del electrodo CD (GEC)")))
    put(grounding, "gec_ac", v.size((t, "Conductor del electrodo CA (GEC)")))
    dc_system = v.choice(
        (t, "Sistema de CD"),
        ["No aterrizado (690-35)", "Aterrizado"],
        default="No aterrizado (690-35)",
    )
    grounding["dc_system"] = DC_SYSTEMS[dc_system.casefold()]
    return request, v.issues


# --- rejections back onto the sheet --------------------------------------------------------------

_SERIES_RULES = {"VOLT-001", "VOLT-002", "VOLT-003", "STR-001", "STR-002", "STR-004"}


def _cell_of(rule: str, subject: str, request: Mapping[str, Any]) -> Key:
    """The cell a rejected fixed value lives in (the inverter when nothing more precise fits)."""
    if rule == "REQ-002":
        return ("Equipos", "Número de cadenas")
    if rule == "REQ-001":
        return ("Equipos", "Cantidad total de módulos")
    if rule in _SERIES_RULES and "n_series" in request:
        return ("Equipos", "Módulos por cadena")
    if rule == "OCP-002":
        if request.get("dc_ocpd_device") == "breaker" or "dc_fuses" not in request:
            return ("Protecciones", "ITM de CD (si elige ITM)")
        return ("Protecciones", "Fusible gPV")
    circuit = "CA" if "INV" in subject else "CD"
    if rule == "CON-002":
        return ("Conductores", f"Calibre de {circuit}")
    if rule == "CON-006":
        return ("Conductores", f"Tubo de {circuit} (EMT)")
    return ("Equipos", "Inversor")


def rejection_issues(
    sizing: SizingResult, request: Mapping[str, Any], cells: Mapping[Key, SheetCell]
) -> list[SheetIssue]:
    """Why no configuration fits, one reason per cell (the first configuration's)."""
    # The engine also tries other module counts on its way; only the count the user asked for
    # (or a whole inverter refused) explains why the sheet cannot be designed.
    total = request.get("module_count_min")
    relevant = [
        r
        for r in sizing.rejected
        if total is None
        or r.config.n_strings == 0
        or r.config.n_strings * r.config.n_series == total
    ]
    # When some string layout was feasible, what failed after it (protection, conductors, rule
    # pack) is the reason; the inverters that never fitted the array are not.
    later = [r for r in relevant if r.stage in ("bos", "rule_pack")]
    found: dict[Key, SheetIssue] = {}
    for rejection in later or relevant or sizing.rejected:
        for issue in rejection.errors:
            key = _cell_of(issue.rule_id, issue.subject, request)
            message = issue.message_es
            if issue.rule_id == "OCP-002" and "Ningún" in message:
                message = message[message.index("Ningún") :]  # drop the engine's preamble
            if key not in found:
                cell = cells.get(key)
                found[key] = SheetIssue(
                    key[0],
                    key[1],
                    f"{issue.rule_id}: {message}",
                    cell.row if cell else None,
                )
    if not found:
        found[("Equipos", "Inversor")] = SheetIssue(
            "Equipos", "Inversor", "Ninguna configuración del catálogo cumple con estos datos."
        )
    return list(found.values())


_EQUIPMENT = {
    ("Equipos", label)
    for label in (
        "Módulo fotovoltaico",
        "Cantidad total de módulos",
        "Número de cadenas",
        "Módulos por cadena",
        "Inversor",
    )
}


def _sizing_issues(
    request: Mapping[str, Any],
    registry: ComponentRegistry,
    cells: Mapping[Key, SheetCell],
    *,
    flagged: set[Key],
) -> list[SheetIssue]:
    """Also size a sheet that has other problems, so the user gets every cell to fix at once.

    Only when the equipment is complete and known; cells already flagged keep their first reason.
    """
    if not request.get("module") or "module_count_min" not in request:
        return []
    if any(cells.get(key) is None for key in flagged) or flagged & _EQUIPMENT:
        return []  # sizing would only restate a problem the sheet already shows
    try:
        sizing = size_pv_system(request_from_mapping(request), registry)
    except (SizingInputError, ValueError):
        return []  # the sheet problems already explain it
    if sizing.selected is not None:
        return []
    return [
        issue
        for issue in rejection_issues(sizing, request, cells)
        if (issue.sheet, issue.label) not in flagged
    ]


# --- workbooks -----------------------------------------------------------------------------------


def write_reviewed_sheet(source: Path, target: Path, issues: Iterable[SheetIssue]) -> Path:
    """The user's workbook with every problem cell in red and its reason as a comment."""
    from openpyxl import load_workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill

    workbook = load_workbook(source)
    red = PatternFill("solid", fgColor=REVIEW_FILL)
    issues = list(issues)
    for issue in issues:
        if issue.row is None or issue.sheet not in workbook.sheetnames:
            continue
        cell = workbook[issue.sheet].cell(issue.row, 2)
        cell.fill = red
        cell.comment = Comment(issue.message_es, "pvsld", width=320, height=120)
    guide = workbook["Instrucciones"] if "Instrucciones" in workbook.sheetnames else None
    if guide is not None:
        row = guide.max_row + 2
        guide.cell(row, 2, f"REVISIÓN: {len(issues)} dato(s) por corregir (celdas en rojo)")
        guide.cell(row, 2).font = Font(name="Arial", bold=True, size=12, color="C00000")
        for offset, issue in enumerate(issues, start=1):
            guide.cell(row + offset, 2, f"• {issue.sheet} › {issue.label}: {issue.message_es}")
            guide.cell(row + offset, 2).font = Font(name="Arial", color="C00000")
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    return target


def refresh_last_calculation(source: Path, fresh: Path, target: Path) -> Path:
    """The user's workbook with column E ("Último cálculo") taken from a freshly written sheet."""
    from openpyxl import load_workbook

    new = load_workbook(fresh)
    computed = {
        (name, str(label).strip()): last
        for name in INPUT_SHEETS
        for label, last in (
            (row[0].value, row[4].value)
            for row in new[name].iter_rows(min_row=FIRST_ROW)
            if row[3].value
        )
        if label
    }
    workbook = load_workbook(source)
    for name in INPUT_SHEETS:
        sheet = workbook[name]
        for row in sheet.iter_rows(min_row=FIRST_ROW):
            label = row[0].value
            if label and row[3].value and (name, str(label).strip()) in computed:
                row[4].value = computed[(name, str(label).strip())]
    workbook.save(target)
    return target


# --- the flow ------------------------------------------------------------------------------------


@dataclass(frozen=True)
class ProResult:
    """Outcome of :func:`design_from_sheet`."""

    ok: bool
    stage: Literal["sheet", "sizing", "validation", "drawing", "done"]
    issues: tuple[SheetIssue, ...] = ()
    design: DesignResult | None = None
    request: dict[str, Any] = field(default_factory=dict)

    @property
    def missing_components(self) -> tuple[str, ...]:
        return tuple(i.missing_component for i in self.issues if i.missing_component)


def design_from_sheet(
    sheet: Path,
    registry: ComponentRegistry,
    folder: Path,
    *,
    use_autocad: Literal["auto", "never"] = "auto",
    autocad: AutocadInfo | None = None,
    runner: Runner = run_process,
    today: date | None = None,
    designer: Callable[..., DesignResult] = design_and_draw,
) -> ProResult:
    """Read, check and design the project sheet; deliverables go to ``folder``.

    The project sheet among the deliverables is the user's own workbook with the last calculation
    refreshed, not a new one.

    Raises:
        ProjectSheetError: the file is not a project sheet.
    """
    cells = read_project_sheet(sheet)
    request, issues = sheet_request(cells, registry, today=today)
    if issues:
        issues += _sizing_issues(
            request, registry, cells, flagged={(i.sheet, i.label) for i in issues}
        )
        return ProResult(False, "sheet", tuple(issues), request=request)
    result = designer(
        request,
        registry,
        folder,
        use_autocad=use_autocad,
        autocad=autocad,
        runner=runner,
        today=today,
        project_sheet=sheet,
    )
    if result.stage == "sizing":
        found = rejection_issues(result.sizing, request, cells)
        return ProResult(False, "sizing", tuple(found), result, request)
    return ProResult(result.ok, result.stage, (), result, request)
