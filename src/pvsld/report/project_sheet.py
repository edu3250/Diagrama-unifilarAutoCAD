"""The professional-mode project sheet (hoja de proyecto): an Excel workbook the user fills in.

Every parameter has a yellow input cell (AUTO, a value or a catalogue dropdown), a gray column
with what the last calculation chose, and a note (owner approval of the prototype, 2026-10-09).
The dropdown lists live on a hidden "Catálogo" sheet built from the catalogue the user has
installed, including the local one. Reading the sheet back is Stage 3.6.2.
"""

from __future__ import annotations

from collections.abc import Iterable
from pathlib import Path
from typing import Any

from pvsld.core.calc import Derived
from pvsld.core.model import PvSystemSpec
from pvsld.core.tables import get_tables

SHEET_VERSION = "1.0"
AUTO = "AUTO"


def _g(value: float) -> str:
    return f"{value:g}"


def _lists(registry: Any, tables: Any) -> dict[str, list[str]]:
    """Dropdown lists: catalogue ids (AUTO first where the calculation may choose) and choices."""
    emt = [f'{d} mm ({t}")' for d, t, _diameter, _area in tables.emt]
    sizes = list(tables.ampacity_cu)
    return {
        "L_Modulos": [c.component_id for c in registry.modules()],
        "L_Inversores": [AUTO] + [c.component_id for c in registry.inverters()],
        "L_Fusibles": [AUTO] + [c.component_id for c in registry.dc_fuses()],
        "L_ITMCD": [AUTO] + [c.component_id for c in registry.dc_breakers()],
        "L_Seccionadores": [AUTO] + [c.component_id for c in registry.dc_switches()],
        "L_ITMCA": [AUTO] + [c.component_id for c in registry.ac_breakers()],
        "L_ITMP": [c.component_id for c in registry.ac_breakers()],
        "L_Cables": [AUTO] + [c.component_id for c in registry.cables()],
        "L_Calibres": [AUTO, *sizes],
        "L_EMT": [AUTO, *emt],
        "L_ProtCadena": ["Fusible gPV", "ITM de CD"],
        "L_Sistema": ["1F-2H", "2F-3H", "3F-4H"],
        "L_Tarifa": [
            "1",
            "1A",
            "1B",
            "1C",
            "1D",
            "1E",
            "1F",
            "DAC",
            "PDBT",
            "GDBT",
            "GDMTO",
            "GDMTH",
        ],
        "L_Regimen": ["Medición neta", "Facturación neta", "Venta total"],
        "L_SiNo": ["No", "Sí"],
        "L_DPS": ["T2", "T1+T2"],
        "L_Electrodo": [
            "Varilla copperweld 3.05 m",
            "Varilla copperweld 2.44 m",
            "Electrodo químico",
        ],
        "L_SistemaCD": ["No aterrizado (690-35)", "Aterrizado"],
        "L_Auto123": [AUTO, "1", "2"],
    }


def _rows(
    spec: PvSystemSpec | None, derived: Derived | None, module_id: str = ""
) -> dict[str, list[Any]]:
    """Sheet rows: section titles, or (label, value, unit, options, last calculation, note)."""

    def first(items: Iterable[Any]) -> Any:
        return next(iter(items), None)

    s = spec
    module = s.modules[0] if s else None
    inv = s.inverters[0] if s else None
    dc_c = first(c for c in s.circuits if c.kind == "pv_source") if s else None
    ac_c = first(c for c in s.circuits if c.kind == "inverter_output") if s else None
    fuse = first(d for d in s.dc_bos.disconnects if d.id.startswith("FUS-")) if s else None
    breaker = first(d for d in s.dc_bos.disconnects if d.id.startswith("DCB-")) if s else None
    switch = first(d for d in s.dc_bos.disconnects if d.id.startswith("DCD-CD")) if s else None
    itm = first(o for o in s.ac_bos.ocpds if o.role == "I1") if s else None
    main = first(s.ac_bos.main_breakers) if s else None
    panel = first(s.ac_bos.panels) if s else None
    dc_spd = first(s.dc_bos.spds) if s else None
    ac_spd = first(s.ac_bos.spds) if s else None

    def race(circuit: Any) -> str:
        if circuit is None or s is None:
            return ""
        owner = (
            circuit
            if circuit.raceway.ref is None
            else next(c for c in s.circuits if c.id == circuit.raceway.ref)
        )
        r = owner.raceway
        if r.type is None or r.trade_size_mm is None:
            return ""
        trade = next(
            (t for d, t, _x, _y in get_tables(s.standards.nom_edition).emt if d == r.trade_size_mm),
            "",
        )
        return f'{_g(r.trade_size_mm)} mm ({trade}")' if trade else f"{_g(r.trade_size_mm)} mm"

    n_modules = sum(x.n_series for x in s.strings) if s else 10
    n_strings = len(s.strings) if s else None
    n_series = s.strings[0].n_series if s else None
    ratio = derived.dc_ac_ratio if derived else None
    return {
        "Proyecto": [
            "Proyecto",
            (
                "Nombre del proyecto",
                s.project.name if s else "",
                "",
                "Texto libre",
                "",
                "Título del plano y de la memoria",
            ),
            ("Titular / propietario", "", "", "Opcional", "", "En blanco = no aparece en el plano"),
            ("Calle y número", "", "", "Opcional"),
            ("Colonia", "", "", "Opcional"),
            ("Municipio", "", "", "Opcional"),
            ("Estado", "", "", "Opcional"),
            ("Código postal", "", "", "Opcional"),
            ("Latitud", "", "°", "Opcional, decimal"),
            ("Longitud", "", "°", "Opcional, decimal"),
            "Servicio CFE",
            ("RPU", "", "", "Opcional (12 dígitos)"),
            ("Número de servicio", "", "", "Opcional"),
            ("Número de medidor", "", "", "Opcional"),
            ("Tarifa", s.utility.tariff if s else "1C", "", "L_Tarifa"),
            ("Sistema de suministro", s.utility.system if s else "2F-3H", "", "L_Sistema"),
            ("Tensión nominal", s.utility.nominal_voltage_v if s else 220, "V", "Número"),
            (
                "Corriente de falla disponible",
                s.utility.available_fault_current_ka if s else 10,
                "kA",
                "Dato de CFE; 10 kA si no se conoce",
                "",
                "Define la capacidad interruptiva mínima de los ITM",
            ),
            ("Régimen", "Medición neta", "", "L_Regimen"),
            "Responsable",
            ("Ingeniero responsable", "", "", "Opcional"),
            ("Cédula profesional", "", "", "Opcional"),
            ("Empresa instaladora", "", "", "Opcional", "", "En blanco: «COMPAÑÍA INSTALADORA»"),
            ("UVIE", "", "", "Opcional"),
            ("Fecha", "", "", "dd/mm/aaaa"),
        ],
        "Equipos": [
            "Arreglo fotovoltaico",
            (
                "Módulo fotovoltaico",
                module_id,
                "",
                "L_Modulos",
                f"{module.manufacturer} {module.model} — {_g(module.pmax_w)} W" if module else "",
            ),
            ("Cantidad total de módulos", n_modules, "pza", "Número"),
            ("Número de cadenas", AUTO, "", "L_Auto123", n_strings or ""),
            (
                "Módulos por cadena",
                AUTO,
                "",
                "AUTO o número",
                n_series or "",
                "Debe dar la cantidad total",
            ),
            ("Optimizadores en todos los módulos", "No", "", "L_SiNo"),
            "Inversor",
            (
                "Inversor",
                AUTO,
                "",
                "L_Inversores",
                f"{inv.manufacturer} {inv.model}" if inv else "",
            ),
            (
                "Relación CD/CA resultante",
                "",
                "",
                "Solo lectura",
                f"{ratio:.2f}" if ratio is not None else "",
            ),
        ],
        "Protecciones": [
            "Caja de protecciones CD",
            (
                "Protección de cadena",
                "ITM de CD" if breaker and not fuse else "Fusible gPV",
                "",
                "L_ProtCadena",
            ),
            (
                "Fusible gPV",
                AUTO,
                "",
                "L_Fusibles",
                (fuse.model or "") if fuse else "—",
                f"Mínimo 1.56 × Isc = {1.5625 * module.isc_a:.1f} A; máximo "
                f"{_g(module.max_series_fuse_a)} A"
                if module
                else "",
            ),
            (
                "ITM de CD (si elige ITM)",
                AUTO,
                "",
                "L_ITMCD",
                (breaker.model or "") if breaker else "—",
            ),
            (
                "Seccionador de la caja",
                AUTO,
                "",
                "L_Seccionadores",
                (switch.model or "") if switch else "—",
                f"{switch.poles} polos, {_g(switch.ie_a)} A a {_g(switch.ue_v)} V"
                if switch
                else "",
            ),
            ("DPS de CD: tipo", dc_spd.spd_type if dc_spd else "T2", "", "L_DPS"),
            ("DPS de CD: corriente nominal In", dc_spd.in_ka if dc_spd else 20, "kA", "Número"),
            "Lado de CA",
            (
                "ITM-1 (salida del inversor)",
                AUTO,
                "",
                "L_ITMCA",
                (itm.model or "") if itm else "",
                f"≥ 1.25 × {_g(inv.iac_max_a)} A = {1.25 * inv.iac_max_a:.1f} A → "
                f"{_g(itm.rating_a)} A, {_g(itm.kaic_ka)} kA"
                if itm and inv
                else "",
            ),
            (
                "ITM-P (principal del servicio)",
                (main.model or "") if main else "",
                "",
                "L_ITMP",
                (main.model or "") if main else "",
                "El existente del servicio",
            ),
            (
                "Barra del centro de carga",
                panel.bus_a if panel else 125,
                "A",
                "Número",
                "",
                "Regla del 120 %: ITM-1 + ITM-P ≤ 1.2 × barra",
            ),
            ("DPS de CA: tipo", ac_spd.spd_type if ac_spd else "T2", "", "L_DPS"),
            ("DPS de CA: Uc", ac_spd.uc_v if ac_spd and ac_spd.uc_v else 175, "V", "Número"),
        ],
        "Conductores": [
            "Recorridos",
            (
                "Longitud cadenas → caja → inversor",
                dc_c.length_m if dc_c else 25,
                "m",
                "Número (un sentido)",
            ),
            (
                "Longitud inversor → ITM-1",
                ac_c.length_m if ac_c else 30,
                "m",
                "Número (un sentido)",
            ),
            (
                "Separación del tubo sobre el techo",
                dc_c.raceway.rooftop_clearance_mm
                if dc_c and dc_c.raceway.rooftop_clearance_mm is not None
                else 25,
                "mm",
                "Número",
                "",
                "Define el incremento de temperatura (tubo al sol)",
            ),
            "Circuito de CD (cadenas)",
            ("Cable de CD", AUTO, "", "L_Cables", dc_c.conductors.insulation if dc_c else ""),
            ("Calibre de CD", AUTO, "", "L_Calibres", dc_c.conductors.size if dc_c else ""),
            ("Tubo de CD (EMT)", AUTO, "", "L_EMT", race(dc_c)),
            (
                "Caída de tensión máxima CD",
                s.standards.vd_limits_pct.dc if s else 1.5,
                "%",
                "Número",
            ),
            "Circuito de CA (salida del inversor)",
            ("Aislamiento de CA", "THHW-LS", "", "Fijo en esta versión"),
            ("Calibre de CA", AUTO, "", "L_Calibres", ac_c.conductors.size if ac_c else ""),
            ("Tubo de CA (EMT)", AUTO, "", "L_EMT", race(ac_c)),
            ("Caída de tensión máxima CA", s.standards.vd_limits_pct.ac if s else 2, "%", "Número"),
        ],
        "Tierra": [
            "Electrodo",
            ("Tipo de electrodo", "Varilla copperweld 3.05 m", "", "L_Electrodo"),
            ("Cantidad de electrodos", s.grounding.electrodes[0].qty if s else 1, "pza", "Número"),
            (
                "Resistencia de diseño",
                s.grounding.design_resistance_ohm if s else 25,
                "Ω",
                "Número",
                "",
                "NOM 250-53",
            ),
            "Conductores de puesta a tierra",
            (
                "Conductor del electrodo CD (GEC)",
                AUTO,
                "",
                "L_Calibres",
                s.grounding.gec_dc if s else "",
            ),
            (
                "Conductor del electrodo CA (GEC)",
                AUTO,
                "",
                "L_Calibres",
                s.grounding.gec_ac if s else "",
            ),
            (
                "Tierra de equipos (EGC)",
                "Igual al conductor del circuito",
                "",
                "Fijo (decisión de diseño)",
            ),
            ("Sistema de CD", "No aterrizado (690-35)", "", "L_SistemaCD"),
        ],
    }


HEADINGS = {
    "Proyecto": "PROYECTO Y SERVICIO CFE",
    "Equipos": "MÓDULOS E INVERSOR",
    "Protecciones": "PROTECCIONES CD Y CA",
    "Conductores": "CONDUCTORES Y CANALIZACIÓN",
    "Tierra": "PUESTA A TIERRA",
}
READ_ONLY = ("Solo lectura", "Fijo en esta versión", "Fijo (decisión de diseño)")


def write_project_sheet(
    path: Path,
    registry: Any,
    spec: PvSystemSpec | None = None,
    derived: Derived | None = None,
    *,
    explanation_es: str = "",
) -> Path:
    """Write the project sheet; with ``spec`` it is prefilled from that design."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation

    font = lambda **k: Font(name="Arial", **k)  # noqa: E731
    head_fill = PatternFill("solid", fgColor="1F4E78")
    sec_fill = PatternFill("solid", fgColor="DDEBF7")
    input_fill, calc_fill = (
        PatternFill("solid", fgColor="FFF2CC"),
        PatternFill("solid", fgColor="EDEDED"),
    )
    thin = Side(style="thin", color="A6A6A6")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="center")
    tables = get_tables(spec.standards.nom_edition if spec else "NOM-001-SEDE-2012")
    lists = _lists(registry, tables)

    wb = Workbook()
    wi = wb.active
    wi.title = "Instrucciones"
    wi.sheet_view.showGridLines = False
    wi.column_dimensions["A"].width = 4
    wi.column_dimensions["B"].width = 110
    title = (
        f"Proyecto: {spec.project.name} (llenada desde el último cálculo)"
        if spec
        else "Hoja nueva: llene los datos o deje AUTO"
    )
    lines = [
        (
            "HOJA DE PROYECTO — DIAGRAMA UNIFILAR FOTOVOLTAICO (modo profesional)",
            font(bold=True, size=16, color="1F4E78"),
        ),
        (title, font(bold=True, size=11)),
        ("", None),
        ("Cómo usarla", font(bold=True, size=12)),
        (
            "1. Llene o cambie las celdas amarillas de las hojas Proyecto, Equipos, Protecciones, Conductores y Tierra.",
            font(),
        ),
        (
            "2. Deje AUTO en lo que quiera que decida el cálculo (inversor, fusible, calibres, tubo...).",
            font(),
        ),
        (
            "3. Guarde el archivo y devuélvalo al plugin:  /unifilar-pro hoja_de_proyecto.xlsx",
            font(bold=True),
        ),
        (
            "4. El plugin valida la hoja. Si algo no cumple la NOM, le regresa este mismo archivo con la celda en rojo y la razón.",
            font(),
        ),
        (
            "5. Si todo cumple: diagrama DWG o DXF, PDF del diagrama, memoria de cálculo (Excel y PDF) y resumen de materiales.",
            font(),
        ),
        ("", None),
        ("Colores", font(bold=True, size=12)),
        (
            "Amarillo: dato de entrada (editable).   Gris: resultado del último cálculo (solo lectura).   Azul: encabezados.",
            font(),
        ),
        ("", None),
        ("Reglas", font(bold=True, size=12)),
        (
            "• Datos personales (titular, dirección, RPU, responsable): si los llena, aparecen en el plano y la memoria; si no, quedan en blanco.",
            font(),
        ),
        (
            "• Equipos: elija de la lista desplegable. Si su equipo no está, escriba el modelo; el plugin le pedirá la hoja técnica (PDF) y lo agregará a su catálogo local.",
            font(),
        ),
        (
            "• Inversor en AUTO: se prefiere una relación CD/CA entre 1.10 y 1.25; el plugin explica la elección.",
            font(),
        ),
        (
            "• El resultado es un borrador: debe revisarlo y firmarlo el ingeniero responsable.",
            font(),
        ),
    ]
    if explanation_es:
        lines += [
            ("", None),
            ("Elección del inversor en el último cálculo", font(bold=True, size=12)),
            (explanation_es, font()),
        ]
    lines += [
        ("", None),
        (
            f"Versión de la hoja: {SHEET_VERSION} · NOM-001-SEDE-2012 · CFE G0100-04",
            font(size=9, color="777777"),
        ),
    ]
    for row, (text, f) in enumerate(lines, 1):
        cell = wi.cell(row, 2, text)
        if f:
            cell.font = f
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    headers = [
        "Parámetro",
        "Valor (edite aquí)",
        "Unidad",
        "Opciones / ayuda",
        "Último cálculo",
        "Nota",
    ]
    widths = (38, 30, 8, 40, 30, 46)
    module_id = ""
    if spec is not None:
        model = spec.modules[0].model
        module_id = next(
            (c.component_id for c in registry.modules() if c.component_id.endswith(model)), model
        )
    for name, rows in _rows(spec, derived, module_id).items():
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws["A1"] = HEADINGS[name]
        ws["A1"].font = font(bold=True, size=14, color="1F4E78")
        ws["A2"] = (
            "Celdas amarillas: datos de entrada. AUTO = lo decide el cálculo. Columna gris: "
            "lo que eligió el último cálculo (solo lectura)."
        )
        ws["A2"].font = font(italic=True, size=9, color="555555")
        for col, text in enumerate(headers, 1):
            cell = ws.cell(4, col, text)
            cell.font, cell.fill, cell.border = font(bold=True, color="FFFFFF"), head_fill, box
        for col, width in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(4, col).column_letter].width = width
        r = 5
        for row in rows:
            if isinstance(row, str):
                for col in range(1, 7):
                    ws.cell(r, col).fill, ws.cell(r, col).border = sec_fill, box
                ws.cell(r, 1, row).font = font(bold=True)
                r += 1
                continue
            label, value, unit, options, last, note = (list(row) + [""] * 6)[:6]
            ws.cell(r, 1, label).font = font()
            cell = ws.cell(r, 2, value)
            if options in READ_ONLY:
                cell.font, cell.fill = font(color="404040"), calc_fill
            else:
                cell.font, cell.fill = font(color="0000FF"), input_fill
            ws.cell(r, 3, unit).font = font(size=9)
            if options in lists:
                dv = DataValidation(
                    type="list",
                    formula1=f"={options}",
                    allow_blank=True,
                    showErrorMessage=True,
                    errorStyle="warning",
                    errorTitle="Valor fuera de la lista",
                    error="Elija de la lista. Para un equipo nuevo escriba su modelo: el "
                    "plugin pedirá su hoja técnica para agregarlo al catálogo local.",
                )
                dv.promptTitle, dv.prompt = label[:32], "Elija de la lista desplegable."
                ws.add_data_validation(dv)
                dv.add(cell)
                help_text = "Lista desplegable" + (
                    " (AUTO o catálogo)" if AUTO in lists[options] else ""
                )
                ws.cell(r, 4, help_text).font = font(size=9)
            else:
                ws.cell(r, 4, options).font = font(size=9)
            last_cell = ws.cell(r, 5, last)
            last_cell.font, last_cell.fill = font(color="404040"), calc_fill
            ws.cell(r, 6, note).font = font(size=9, color="555555")
            for col in range(1, 7):
                ws.cell(r, col).border, ws.cell(r, col).alignment = box, wrap
            r += 1
        ws.freeze_panes = "A5"

    wc = wb.create_sheet("Catálogo")
    for col, (name, items) in enumerate(lists.items(), 1):
        wc.cell(1, col, name).font = font(bold=True)
        for row, item in enumerate(items or [""], 2):
            wc.cell(row, col, item)
        letter = wc.cell(1, col).column_letter
        wb.defined_names[name] = DefinedName(
            name, attr_text=f"'Catálogo'!${letter}$2:${letter}${max(len(items), 1) + 1}"
        )
    wc.sheet_state = "hidden"
    wb.active = 0
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
